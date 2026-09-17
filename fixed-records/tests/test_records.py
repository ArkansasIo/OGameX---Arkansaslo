#!/usr/bin/env python3
import sqlite3, struct, subprocess, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'tools'))
import record_tool
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
subprocess.run([sys.executable,str(ROOT/'tools/build_excel.py')],check=True)
subprocess.run([sys.executable,str(ROOT/'tools/record_tool.py'),'generate','--output',str(ROOT/'data/records.bin')],check=True)
data=(ROOT/'data/records.bin').read_bytes()
assert len(data)==60 and len(data)%20==0
records=record_tool.validate(data)
assert [r['record_id'] for r in records]==[1001,2001,3001]
assert record_tool.crc16(data[:20]) == 0x6DA4
assert struct.unpack('<I',data[4:8])[0]==1001
conn=sqlite3.connect(':memory:'); conn.executescript((ROOT/'sql/schema.sql').read_text())
rows=conn.execute('select record_id,value,quantity from fixed_records order by record_id').fetchall()
assert rows==[(1001,2500,1),(2001,90000,12),(3001,450,320)]
wb=load_workbook(ROOT/'excel_template.xlsx',data_only=False); ws=wb['Records']
assert ws.max_row==4 and ws['E2'].value==1001 and ws['L2'].value.startswith('=AND')
print('PASS: binary, CRC, SQLite, and Excel artifacts agree')
